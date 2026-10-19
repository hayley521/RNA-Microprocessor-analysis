#!/usr/bin/env python3
"""Export publication-ready bulk RNA-seq processed tables.

The script is deliberately strict: source columns and Excel worksheet headers must
match the explicitly reviewed layouts below.  Unknown or missing columns stop the
run instead of being assigned by position.

The DICER1 featureCounts/rMATS sources live on a different server from the other
inputs.  Use repeated ``--only`` options to run the relevant tasks on each server,
then copy the generated files into the same submission staging directory.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Iterable, Iterator, Sequence


ANNOTATION_COLUMNS = ("Chr", "Start", "End", "Strand", "Length")
TPM_ANNOTATION_COLUMNS = ("ENSEMBL", "Gene_symbol", "Gene_type")
EVENT_TYPES = ("A3SS", "A5SS", "MXE", "RI", "SE")
TASKS = (
    "counts-gr",
    "counts-dicer",
    "counts-basal",
    "tpm",
    "de",
    "rmats-gr",
    "rmats-dicer",
)


# source alias -> standard sample_id in raw_data_manifest.tsv
GR_COUNT_MAP = {
    **{
        f"L{rep}_G{time_h}": f"bulk_DGCR8_{time_h}h_rep{rep}"
        for time_h in (0, 4, 8, 12)
        for rep in (1, 2)
    },
    **{
        f"L{rep}_R{time_h}": f"bulk_DROSHA_{time_h}h_rep{rep}"
        for time_h in (0, 4, 8, 12)
        for rep in (1, 2)
    },
}

DICER_COUNT_MAP = {
    "I1_I13-0h-1": "bulk_DICER1_0h_rep1",
    "I3_I13-0h-2": "bulk_DICER1_0h_rep2",
    "I1_I13-4h-1": "bulk_DICER1_4h_rep1",
    "I3_I13-4h-2": "bulk_DICER1_4h_rep2",
    "I2_I13-8h-1": "bulk_DICER1_8h_rep1",
    "I4_I13-8h-2": "bulk_DICER1_8h_rep2",
    "I2_I13-12h-1": "bulk_DICER1_12h_rep1",
    "I4_I13-12h-2": "bulk_DICER1_12h_rep2",
}
DICER_COUNT_EXCLUDED = {
    "I1_I13-1h-1",
    "I1_I13-2h-1",
    "I2_I13-24h-1",
    "I2_I13-72h-1",
    "I3_I13-1h-2",
    "I3_I13-2h-2",
    "I4_I13-24h-2",
    "I4_I13-72h-2",
}

BASAL_COUNT_MAP = {
    f"CRGI_{source}-{rep}": f"bulk_basal_{target}_rep{rep}"
    for target, source in (
        ("WT", "WT"),
        ("DGCR8", "G13"),
        ("DROSHA", "R1"),
        ("DICER1", "I13"),
    )
    for rep in (1, 2)
}

GR_TPM_MAP = {
    **{
        f"G_{time_h}h_{rep}": f"bulk_DGCR8_{time_h}h_rep{rep}"
        for time_h in (0, 4, 8, 12)
        for rep in (1, 2)
    },
    **{
        f"R_{time_h}h_{rep}": f"bulk_DROSHA_{time_h}h_rep{rep}"
        for time_h in (0, 4, 8, 12)
        for rep in (1, 2)
    },
}
DICER_TPM_MAP = {
    f"I_{time_h}h_{rep}": f"bulk_DICER1_{time_h}h_rep{rep}"
    for time_h in (0, 4, 8, 12)
    for rep in (1, 2)
}
DICER_TPM_EXCLUDED = {
    f"I_{time_h}h_{rep}"
    for time_h in (1, 2, 24, 72)
    for rep in (1, 2)
}


@dataclass(frozen=True)
class CountSpec:
    task: str
    source_arg: str
    output_name: str
    sample_map: dict[str, str]
    excluded_aliases: set[str]


@dataclass(frozen=True)
class TpmSpec:
    sheet: str
    output_name: str
    sample_map: dict[str, str]
    excluded_columns: set[str]


@dataclass(frozen=True)
class DeSpec:
    target: str
    source_arg: str
    sheet: str
    time_h: int
    prefix: str

    @property
    def output_name(self) -> str:
        return f"bulk_RNAseq_{self.target}_{self.time_h}h_vs_0h_DE.tsv"

    @property
    def expected_header(self) -> tuple[str, ...]:
        return (
            "Gene_symbol",
            "Gene_type",
            f"{self.prefix}_log2FoldChange",
            f"{self.prefix}_pvalue",
            f"{self.prefix}_padj",
        )


@dataclass(frozen=True)
class ExportRecord:
    task: str
    source: str
    output: str
    data_rows: int
    columns: int
    sha256: str


COUNT_SPECS = (
    CountSpec(
        "counts-gr",
        "gr_counts",
        "bulk_RNAseq_DGCR8_DROSHA_raw_counts.tsv",
        GR_COUNT_MAP,
        set(),
    ),
    CountSpec(
        "counts-dicer",
        "dicer_counts",
        "bulk_RNAseq_DICER1_raw_counts.tsv",
        DICER_COUNT_MAP,
        DICER_COUNT_EXCLUDED,
    ),
    CountSpec(
        "counts-basal",
        "basal_counts",
        "bulk_RNAseq_basal_WT_AID_raw_counts.tsv",
        BASAL_COUNT_MAP,
        set(),
    ),
)

TPM_SPECS = (
    TpmSpec(
        "Sheet1",
        "bulk_RNAseq_DGCR8_DROSHA_TPM.tsv",
        GR_TPM_MAP,
        set(),
    ),
    TpmSpec(
        "Sheet2",
        "bulk_RNAseq_DICER1_TPM.tsv",
        DICER_TPM_MAP,
        DICER_TPM_EXCLUDED,
    ),
)

DE_SPECS = tuple(
    DeSpec(target, source_arg, f"{prefix}{time_h}", time_h, f"{prefix}{time_h}")
    for target, source_arg, prefix in (
        ("DGCR8", "dgcr8_de_workbook", "G"),
        ("DROSHA", "drosha_de_workbook", "R"),
        ("DICER1", "dicer1_de_workbook", "I"),
    )
    for time_h in (4, 8, 12)
)


class ExportError(RuntimeError):
    """Raised when a reviewed source layout does not validate."""


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Clean and rename final bulk RNA-seq count, TPM, DE, and rMATS "
            "tables for manuscript submission."
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        required=True,
        help="Destination directory for submission-ready TSV files.",
    )
    parser.add_argument(
        "--gr-counts",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/RNA-seq/rawCounts_O.tsv"),
        help="DGCR8/DROSHA featureCounts rawCounts_O.tsv.",
    )
    parser.add_argument(
        "--dicer-counts",
        type=Path,
        default=Path("/home/xiegang/data/WKW/250531/rawCounts_O.tsv"),
        help="DICER1 featureCounts rawCounts_O.tsv.",
    )
    parser.add_argument(
        "--basal-counts",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/260226/rawCounts_O.tsv"),
        help="WT/AID basal featureCounts rawCounts_O.tsv.",
    )
    parser.add_argument(
        "--tpm-workbook",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/TPM.xlsx"),
        help="Workbook containing reviewed Sheet1 (G/R) and Sheet2 (DICER1).",
    )
    parser.add_argument(
        "--dgcr8-de-workbook",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/DGCR8-AID.xlsx"),
    )
    parser.add_argument(
        "--drosha-de-workbook",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/DROSHA-AID.xlsx"),
    )
    parser.add_argument(
        "--dicer1-de-workbook",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/DICER1-AID.xlsx"),
    )
    parser.add_argument(
        "--gr-rmats-root",
        type=Path,
        default=Path("/home/xinhuil/lyw-GROseq/RNA-seq"),
        help="Directory containing rmats_output_G4 ... rmats_output_R12.",
    )
    parser.add_argument(
        "--dicer-rmats-root",
        type=Path,
        default=Path("/home/xiegang/data/WKW/250531/rmats_results"),
        help="Directory containing DICER1 4h_vs_0h, 8h_vs_0h, and 12h_vs_0h.",
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        help=(
            "Optional raw_data_manifest.tsv. When supplied, every emitted standard "
            "sample ID must exist in its bulk RNA-seq rows."
        ),
    )
    parser.add_argument(
        "--only",
        action="append",
        choices=TASKS,
        metavar="TASK",
        help=(
            "Run only this task; repeat for multiple tasks. Omit to run all tasks. "
            f"Choices: {', '.join(TASKS)}"
        ),
    )
    annotation_group = parser.add_mutually_exclusive_group()
    annotation_group.add_argument(
        "--keep-annotations",
        dest="keep_annotations",
        action="store_true",
        default=True,
        help="Retain available featureCounts Chr/Start/End/Strand/Length columns.",
    )
    annotation_group.add_argument(
        "--drop-annotations",
        dest="keep_annotations",
        action="store_false",
        help="Emit Geneid and sample count columns only.",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate selected inputs and print the export plan without writing.",
    )
    mode.add_argument(
        "--validate-only",
        action="store_true",
        help="Validate selected inputs and already exported outputs without writing.",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace existing planned outputs during a real export.",
    )
    return parser.parse_args(argv)


def _selected_tasks(args: argparse.Namespace) -> tuple[str, ...]:
    if not args.only:
        return TASKS
    requested = set(args.only)
    return tuple(task for task in TASKS if task in requested)


def _require_file(path: Path, label: str) -> None:
    if not path.is_file():
        raise ExportError(f"{label} is missing or is not a regular file: {path}")


def _trim_trailing_empty(values: Sequence[object]) -> tuple[object, ...]:
    result = list(values)
    while result and result[-1] is None:
        result.pop()
    return tuple(result)


def _is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _display_value(value: object) -> object:
    return "" if value is None else value


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _atomic_path(path: Path) -> Path:
    return path.with_name(f".{path.name}.tmp.{os.getpid()}")


def _write_tsv_atomic(
    output: Path,
    header: Sequence[object],
    rows: Iterable[Sequence[object]],
) -> int:
    temporary = _atomic_path(output)
    data_rows = 0
    try:
        with temporary.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
            writer.writerow([_display_value(value) for value in header])
            for row in rows:
                writer.writerow([_display_value(value) for value in row])
                data_rows += 1
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    return data_rows


def _copy_atomic(source: Path, output: Path) -> None:
    temporary = _atomic_path(output)
    try:
        shutil.copyfile(source, temporary)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()


def _read_featurecounts_header(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for line in handle:
            if not line.strip() or line.startswith("#"):
                continue
            return next(csv.reader([line], delimiter="\t"))
    raise ExportError(f"No featureCounts header found in {path}")


def _featurecounts_rows(path: Path) -> Iterator[list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(
            (line for line in handle if line.strip() and not line.startswith("#")),
            delimiter="\t",
        )
        try:
            next(reader)
        except StopIteration as exc:
            raise ExportError(f"No featureCounts header found in {path}") from exc
        yield from reader


def _column_matches_alias(column: str, alias: str) -> bool:
    normalized = column.replace("\\", "/")
    path = PurePosixPath(normalized)
    name = path.name
    parent = path.parent.name
    return (
        column == alias
        or name == alias
        or parent == alias
        or name.startswith(f"{alias}_")
    )


def _resolve_featurecounts_columns(
    header: Sequence[str],
    sample_map: dict[str, str],
    excluded_aliases: set[str],
) -> tuple[list[int], list[str], list[int], list[str]]:
    if not header or header[0] != "Geneid":
        raise ExportError("featureCounts first column must be exactly 'Geneid'")
    if len(set(header)) != len(header):
        raise ExportError("Duplicate featureCounts column names were found")

    annotation_indexes = [
        header.index(name) for name in ANNOTATION_COLUMNS if name in header
    ]
    metadata = {"Geneid", *ANNOTATION_COLUMNS}
    sample_columns = [(index, name) for index, name in enumerate(header) if name not in metadata]
    selected_by_alias: dict[str, int] = {}
    known_aliases = tuple(sample_map) + tuple(sorted(excluded_aliases))

    for index, column in sample_columns:
        matches = [alias for alias in known_aliases if _column_matches_alias(column, alias)]
        if len(matches) != 1:
            detail = "no reviewed alias" if not matches else f"multiple aliases: {matches}"
            raise ExportError(f"Unrecognized featureCounts sample column {column!r} ({detail})")
        alias = matches[0]
        if alias in sample_map:
            if alias in selected_by_alias:
                raise ExportError(f"Multiple source columns matched alias {alias!r}")
            selected_by_alias[alias] = index

    missing = [alias for alias in sample_map if alias not in selected_by_alias]
    if missing:
        raise ExportError(f"Missing reviewed featureCounts samples: {', '.join(missing)}")

    selected_indexes = [selected_by_alias[alias] for alias in sample_map]
    standard_ids = list(sample_map.values())
    return annotation_indexes, [header[i] for i in annotation_indexes], selected_indexes, standard_ids


def validate_count_input(path: Path, spec: CountSpec) -> tuple[list[int], list[str], list[int], list[str]]:
    _require_file(path, spec.task)
    header = _read_featurecounts_header(path)
    return _resolve_featurecounts_columns(header, spec.sample_map, spec.excluded_aliases)


def export_counts(
    path: Path,
    output: Path,
    spec: CountSpec,
    keep_annotations: bool,
) -> ExportRecord:
    header = _read_featurecounts_header(path)
    annotation_i, annotation_names, sample_i, standard_ids = _resolve_featurecounts_columns(
        header, spec.sample_map, spec.excluded_aliases
    )
    if not keep_annotations:
        annotation_i, annotation_names = [], []
    output_header = ["Geneid", *annotation_names, *standard_ids]

    def cleaned_rows() -> Iterator[list[object]]:
        for line_number, row in enumerate(_featurecounts_rows(path), start=2):
            if len(row) != len(header):
                raise ExportError(
                    f"{path}:{line_number}: expected {len(header)} fields, found {len(row)}"
                )
            if not row[0]:
                raise ExportError(f"{path}:{line_number}: empty Geneid")
            for index in sample_i:
                try:
                    float(row[index])
                except ValueError as exc:
                    raise ExportError(
                        f"{path}:{line_number}: non-numeric count {row[index]!r} in {header[index]!r}"
                    ) from exc
            yield [row[0], *(row[i] for i in annotation_i), *(row[i] for i in sample_i)]

    rows = _write_tsv_atomic(output, output_header, cleaned_rows())
    return ExportRecord(spec.task, str(path), output.name, rows, len(output_header), _sha256(output))


def _load_openpyxl():
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise ExportError(
            "openpyxl is required for Excel inputs. Install it in the active conda environment."
        ) from exc
    return openpyxl


def _worksheet_header(worksheet) -> tuple[object, ...]:
    iterator = worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
    try:
        return _trim_trailing_empty(next(iterator))
    except StopIteration as exc:
        raise ExportError(f"Worksheet {worksheet.title!r} is empty") from exc


def _resolve_tpm_columns(
    header: Sequence[object], spec: TpmSpec
) -> tuple[list[int], list[int], list[str]]:
    if any(not isinstance(value, str) for value in header):
        raise ExportError(f"TPM {spec.sheet} contains a blank or non-text header: {header}")
    text_header = list(header)  # type: ignore[arg-type]
    if len(set(text_header)) != len(text_header):
        raise ExportError(f"TPM {spec.sheet} contains duplicate headers")
    missing_annotations = [name for name in TPM_ANNOTATION_COLUMNS if name not in text_header]
    if missing_annotations:
        raise ExportError(
            f"TPM {spec.sheet} is missing annotation columns: {', '.join(missing_annotations)}"
        )
    annotation_i = [text_header.index(name) for name in TPM_ANNOTATION_COLUMNS]
    known = set(TPM_ANNOTATION_COLUMNS) | set(spec.sample_map) | spec.excluded_columns
    unknown = [name for name in text_header if name not in known]
    if unknown:
        raise ExportError(f"TPM {spec.sheet} has unreviewed columns: {unknown}")
    missing_samples = [name for name in spec.sample_map if name not in text_header]
    if missing_samples:
        raise ExportError(f"TPM {spec.sheet} is missing samples: {missing_samples}")
    sample_i = [text_header.index(name) for name in spec.sample_map]
    return annotation_i, sample_i, list(spec.sample_map.values())


def validate_tpm_input(path: Path) -> None:
    _require_file(path, "TPM workbook")
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for spec in TPM_SPECS:
            if spec.sheet not in workbook.sheetnames:
                raise ExportError(f"TPM workbook is missing sheet {spec.sheet!r}")
            header = _worksheet_header(workbook[spec.sheet])
            _resolve_tpm_columns(header, spec)
    finally:
        workbook.close()


def export_tpm(path: Path, output_dir: Path) -> list[ExportRecord]:
    openpyxl = _load_openpyxl()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records: list[ExportRecord] = []
    try:
        for spec in TPM_SPECS:
            worksheet = workbook[spec.sheet]
            header = _worksheet_header(worksheet)
            annotation_i, sample_i, standard_ids = _resolve_tpm_columns(header, spec)
            output_header = [*TPM_ANNOTATION_COLUMNS, *standard_ids]

            def cleaned_rows() -> Iterator[list[object]]:
                nonempty_by_sample = [0] * len(sample_i)
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    values = tuple(row[: len(header)])
                    if all(_is_empty(value) for value in values):
                        continue
                    for position, index in enumerate(sample_i):
                        if index < len(values) and not _is_empty(values[index]):
                            nonempty_by_sample[position] += 1
                    yield [
                        *(values[i] if i < len(values) else None for i in annotation_i),
                        *(values[i] if i < len(values) else None for i in sample_i),
                    ]
                empty_samples = [
                    standard_ids[i]
                    for i, count in enumerate(nonempty_by_sample)
                    if count == 0
                ]
                if empty_samples:
                    raise ExportError(
                        f"TPM {spec.sheet} has no values for samples: {empty_samples}"
                    )

            output = output_dir / spec.output_name
            rows = _write_tsv_atomic(output, output_header, cleaned_rows())
            records.append(
                ExportRecord("tpm", f"{path}::{spec.sheet}", output.name, rows, len(output_header), _sha256(output))
            )
    finally:
        workbook.close()
    return records


def validate_de_inputs(args: argparse.Namespace) -> None:
    openpyxl = _load_openpyxl()
    grouped: dict[Path, list[DeSpec]] = {}
    for spec in DE_SPECS:
        grouped.setdefault(getattr(args, spec.source_arg), []).append(spec)
    for path, specs in grouped.items():
        _require_file(path, f"{specs[0].target} DE workbook")
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for spec in specs:
                if spec.sheet not in workbook.sheetnames:
                    raise ExportError(f"{path} is missing DE sheet {spec.sheet!r}")
                actual = _worksheet_header(workbook[spec.sheet])
                if actual != spec.expected_header:
                    raise ExportError(
                        f"{path}::{spec.sheet} header mismatch; expected "
                        f"{spec.expected_header}, found {actual}"
                    )
        finally:
            workbook.close()


def export_de(args: argparse.Namespace, output_dir: Path) -> list[ExportRecord]:
    openpyxl = _load_openpyxl()
    records: list[ExportRecord] = []
    grouped: dict[Path, list[DeSpec]] = {}
    for spec in DE_SPECS:
        grouped.setdefault(getattr(args, spec.source_arg), []).append(spec)

    for path, specs in grouped.items():
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for spec in specs:
                worksheet = workbook[spec.sheet]
                header = _worksheet_header(worksheet)

                def cleaned_rows() -> Iterator[tuple[object, ...]]:
                    for row in worksheet.iter_rows(min_row=2, values_only=True):
                        values = tuple(row[: len(header)])
                        if all(_is_empty(value) for value in values):
                            continue
                        yield values

                output = output_dir / spec.output_name
                rows = _write_tsv_atomic(output, header, cleaned_rows())
                records.append(
                    ExportRecord(
                        "de",
                        f"{path}::{spec.sheet}",
                        output.name,
                        rows,
                        len(header),
                        _sha256(output),
                    )
                )
        finally:
            workbook.close()
    return records


RMATS_REQUIRED_COLUMNS = {
    "ID",
    "GeneID",
    "geneSymbol",
    "chr",
    "strand",
    "IJC_SAMPLE_1",
    "SJC_SAMPLE_1",
    "IJC_SAMPLE_2",
    "SJC_SAMPLE_2",
    "IncFormLen",
    "SkipFormLen",
    "PValue",
    "FDR",
    "IncLevel1",
    "IncLevel2",
    "IncLevelDifference",
}


def _rmats_specs(task: str, args: argparse.Namespace) -> Iterator[tuple[Path, str]]:
    if task == "rmats-gr":
        root = args.gr_rmats_root
        for code, target in (("G", "DGCR8"), ("R", "DROSHA")):
            for time_h in (4, 8, 12):
                source_dir = root / f"rmats_output_{code}{time_h}"
                for event in EVENT_TYPES:
                    source = source_dir / f"{event}.MATS.JCEC.txt"
                    output = f"bulk_RNAseq_{target}_{time_h}h_vs_0h_rMATS_{event}_JCEC.tsv"
                    yield source, output
    elif task == "rmats-dicer":
        root = args.dicer_rmats_root
        for time_h in (4, 8, 12):
            source_dir = root / f"{time_h}h_vs_0h"
            for event in EVENT_TYPES:
                source = source_dir / f"{event}.MATS.JCEC.txt"
                output = f"bulk_RNAseq_DICER1_{time_h}h_vs_0h_rMATS_{event}_JCEC.tsv"
                yield source, output
    else:
        raise AssertionError(f"Unexpected rMATS task: {task}")


def _rmats_header(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        first = handle.readline().rstrip("\r\n")
    if not first:
        raise ExportError(f"Empty rMATS table: {path}")
    return first.split("\t")


def validate_rmats_input(path: Path) -> list[str]:
    _require_file(path, "rMATS JCEC table")
    header = _rmats_header(path)
    missing = sorted(RMATS_REQUIRED_COLUMNS - set(header))
    if missing:
        raise ExportError(f"{path} is missing rMATS columns: {missing}")
    if header.count("ID") < 1:
        raise ExportError(f"{path} has no rMATS ID column")
    return header


def _line_count(path: Path) -> int:
    with path.open("rb") as handle:
        return sum(1 for _ in handle)


def export_rmats(task: str, args: argparse.Namespace, output_dir: Path) -> list[ExportRecord]:
    records: list[ExportRecord] = []
    for source, output_name in _rmats_specs(task, args):
        header = validate_rmats_input(source)
        output = output_dir / output_name
        _copy_atomic(source, output)
        records.append(
            ExportRecord(
                task,
                str(source),
                output.name,
                max(_line_count(output) - 1, 0),
                len(header),
                _sha256(output),
            )
        )
    return records


def validate_manifest(path: Path, required_ids: set[str]) -> None:
    _require_file(path, "raw-data manifest")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        required_columns = {"sample_id", "assay"}
        if not reader.fieldnames or not required_columns.issubset(reader.fieldnames):
            raise ExportError(f"Manifest lacks required columns: {sorted(required_columns)}")
        manifest_ids = {
            row["sample_id"]
            for row in reader
            if row["assay"] == "bulk_RNA-seq"
        }
    missing = sorted(required_ids - manifest_ids)
    if missing:
        raise ExportError(f"Standard sample IDs missing from manifest: {missing}")


def expected_output_names(tasks: Sequence[str]) -> list[str]:
    names: list[str] = []
    selected = set(tasks)
    names.extend(spec.output_name for spec in COUNT_SPECS if spec.task in selected)
    if "tpm" in selected:
        names.extend(spec.output_name for spec in TPM_SPECS)
    if "de" in selected:
        names.extend(spec.output_name for spec in DE_SPECS)
    # rMATS names are added separately because they require argparse source roots.
    return names


def validate_inputs(args: argparse.Namespace, tasks: Sequence[str]) -> None:
    selected = set(tasks)
    required_ids: set[str] = set()
    for spec in COUNT_SPECS:
        if spec.task in selected:
            validate_count_input(getattr(args, spec.source_arg), spec)
            required_ids.update(spec.sample_map.values())
    if "tpm" in selected:
        validate_tpm_input(args.tpm_workbook)
        for spec in TPM_SPECS:
            required_ids.update(spec.sample_map.values())
    if "de" in selected:
        validate_de_inputs(args)
    for task in ("rmats-gr", "rmats-dicer"):
        if task in selected:
            for source, _ in _rmats_specs(task, args):
                validate_rmats_input(source)
    if args.manifest:
        validate_manifest(args.manifest, required_ids)


def _planned_outputs(args: argparse.Namespace, tasks: Sequence[str]) -> list[Path]:
    names = expected_output_names(tasks)
    for task in ("rmats-gr", "rmats-dicer"):
        if task in tasks:
            names.extend(output_name for _, output_name in _rmats_specs(task, args))
    return [args.output_dir / name for name in names]


def _validate_simple_tsv(path: Path, expected_header: Sequence[str]) -> int:
    _require_file(path, "exported TSV")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        try:
            actual_header = next(reader)
        except StopIteration as exc:
            raise ExportError(f"Exported TSV is empty: {path}") from exc
        if actual_header != list(expected_header):
            raise ExportError(
                f"{path} header mismatch; expected {list(expected_header)}, found {actual_header}"
            )
        rows = 0
        for line_number, row in enumerate(reader, start=2):
            if not row or all(not value.strip() for value in row):
                raise ExportError(f"{path}:{line_number}: fully empty row remains")
            if len(row) != len(actual_header):
                raise ExportError(
                    f"{path}:{line_number}: expected {len(actual_header)} fields, found {len(row)}"
                )
            rows += 1
    return rows


def validate_outputs(args: argparse.Namespace, tasks: Sequence[str]) -> None:
    selected = set(tasks)
    for spec in COUNT_SPECS:
        if spec.task not in selected:
            continue
        annotations = list(ANNOTATION_COLUMNS) if args.keep_annotations else []
        expected = ["Geneid", *annotations, *spec.sample_map.values()]
        _validate_simple_tsv(args.output_dir / spec.output_name, expected)
    if "tpm" in selected:
        for spec in TPM_SPECS:
            expected = [*TPM_ANNOTATION_COLUMNS, *spec.sample_map.values()]
            _validate_simple_tsv(args.output_dir / spec.output_name, expected)
    if "de" in selected:
        for spec in DE_SPECS:
            _validate_simple_tsv(args.output_dir / spec.output_name, spec.expected_header)
    for task in ("rmats-gr", "rmats-dicer"):
        if task not in selected:
            continue
        for source, output_name in _rmats_specs(task, args):
            output = args.output_dir / output_name
            validate_rmats_input(output)
            if _sha256(source) != _sha256(output):
                raise ExportError(f"rMATS copy differs from source: {output}")


def write_export_manifest(output_dir: Path, records: Sequence[ExportRecord]) -> Path:
    output = output_dir / "bulk_processed_export_manifest.tsv"
    header = ("task", "source", "output", "data_rows", "columns", "sha256")
    rows = (
        (record.task, record.source, record.output, record.data_rows, record.columns, record.sha256)
        for record in records
    )
    _write_tsv_atomic(output, header, rows)
    return output


def run_export(args: argparse.Namespace, tasks: Sequence[str]) -> list[ExportRecord]:
    args.output_dir.mkdir(parents=True, exist_ok=True)
    records: list[ExportRecord] = []
    selected = set(tasks)
    for spec in COUNT_SPECS:
        if spec.task in selected:
            source = getattr(args, spec.source_arg)
            output = args.output_dir / spec.output_name
            records.append(export_counts(source, output, spec, args.keep_annotations))
    if "tpm" in selected:
        records.extend(export_tpm(args.tpm_workbook, args.output_dir))
    if "de" in selected:
        records.extend(export_de(args, args.output_dir))
    for task in ("rmats-gr", "rmats-dicer"):
        if task in selected:
            records.extend(export_rmats(task, args, args.output_dir))
    return records


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    tasks = _selected_tasks(args)
    try:
        validate_inputs(args, tasks)
        planned = _planned_outputs(args, tasks)
        print(f"Validated inputs for {len(tasks)} task(s): {', '.join(tasks)}")
        for path in planned:
            print(f"PLAN\t{path}")

        if args.dry_run:
            print(f"Dry run complete; {len(planned)} files planned and nothing written.")
            return 0
        if args.validate_only:
            validate_outputs(args, tasks)
            print(f"Validated {len(planned)} existing output files.")
            return 0

        manifest_path = args.output_dir / "bulk_processed_export_manifest.tsv"
        collisions = [path for path in [*planned, manifest_path] if path.exists()]
        if collisions and not args.overwrite:
            preview = "\n  ".join(str(path) for path in collisions)
            raise ExportError(
                "Planned outputs already exist; use --overwrite to replace them:\n  " + preview
            )

        records = run_export(args, tasks)
        validate_outputs(args, tasks)
        manifest = write_export_manifest(args.output_dir, records)
        print(f"Exported and validated {len(records)} files.")
        print(f"Export manifest: {manifest}")
        return 0
    except (ExportError, OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
